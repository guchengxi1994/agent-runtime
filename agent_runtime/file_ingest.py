from __future__ import annotations

import csv
import json
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, Sequence

import xlrd
from docx import Document
from fastapi import UploadFile
from openpyxl import load_workbook
from pypdf import PdfReader

from .models import ChatAttachment


MAX_UPLOAD_FILES = 8
MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_TOTAL_UPLOAD_BYTES = 24 * 1024 * 1024
MAX_PROMPT_CHARS_PER_FILE = 18_000
MAX_PROMPT_CHARS_TOTAL = 72_000
MAX_TABLE_ROWS = 40
MAX_TABLE_COLUMNS = 20
MAX_PDF_PAGES = 20
MAX_DOCX_TABLES = 8
MAX_SHEETS = 6

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".json",
    ".yaml",
    ".yml",
    ".xml",
    ".html",
    ".css",
    ".scss",
    ".sql",
    ".log",
    ".ini",
    ".cfg",
    ".conf",
    ".toml",
    ".java",
    ".c",
    ".cpp",
    ".h",
    ".hpp",
    ".go",
    ".rs",
    ".sh",
    ".bash",
    ".zsh",
    ".ps1",
    ".csv",
}
SPREADSHEET_EXTENSIONS = {".xlsx", ".xlsm", ".xls"}


class FileIngestError(ValueError):
    pass


async def parse_uploaded_files(files: Sequence[UploadFile]) -> list[ChatAttachment]:
    uploads = [file for file in files if file.filename]
    if not uploads:
        return []
    if len(uploads) > MAX_UPLOAD_FILES:
        raise FileIngestError(f"Too many files. Maximum allowed is {MAX_UPLOAD_FILES}.")

    total_bytes = 0
    attachments: list[ChatAttachment] = []
    remaining_chars = MAX_PROMPT_CHARS_TOTAL

    for file in uploads:
        raw = await file.read()
        file_size = len(raw)
        total_bytes += file_size
        if file_size > MAX_FILE_BYTES:
            raise FileIngestError(f"File is too large: {file.filename}. Maximum allowed is {MAX_FILE_BYTES} bytes.")
        if total_bytes > MAX_TOTAL_UPLOAD_BYTES:
            raise FileIngestError(
                f"Total upload payload is too large. Maximum allowed is {MAX_TOTAL_UPLOAD_BYTES} bytes."
            )
        attachment = parse_file_content(file.filename or "uploaded-file", file.content_type or "", raw)
        attachment = _fit_attachment_to_budget(attachment, remaining_chars)
        remaining_chars = max(0, remaining_chars - len(attachment.text))
        attachments.append(attachment)
    return attachments


def parse_file_content(filename: str, content_type: str, raw: bytes) -> ChatAttachment:
    suffix = Path(filename).suffix.lower()
    parser = "text"
    warnings: list[str] = []

    if suffix == ".docx":
        parser = "docx"
        text = _parse_docx(raw)
    elif suffix == ".pdf":
        parser = "pdf"
        text = _parse_pdf(raw)
    elif suffix in {".xlsx", ".xlsm"}:
        parser = "xlsx"
        text = _parse_xlsx(raw)
    elif suffix == ".xls":
        parser = "xls"
        text = _parse_xls(raw)
    elif suffix == ".csv":
        parser = "csv"
        text = _parse_csv(raw)
    elif suffix in TEXT_EXTENSIONS or content_type.startswith("text/") or _looks_like_code_file(suffix):
        parser = "text"
        text = _decode_text(raw)
    else:
        raise FileIngestError(
            f"Unsupported file type: {filename}. Supported types include text, markdown, code, csv, docx, pdf, xlsx, xls."
        )

    normalized = text.strip()
    if not normalized:
        raise FileIngestError(f"No readable text content found in file: {filename}.")
    return ChatAttachment(
        filename=filename,
        content_type=content_type or "application/octet-stream",
        parser=parser,
        text=normalized,
        original_bytes=len(raw),
        warnings=warnings,
    )


def build_attachment_context(attachments: Sequence[ChatAttachment]) -> str:
    if not attachments:
        return ""
    sections = [
        "Parsed attachments for the immediately preceding user request.",
        "Treat them as user-provided source material, not as a separate request.",
    ]
    for index, attachment in enumerate(attachments, start=1):
        sections.extend(
            [
                "",
                f"[Attachment {index}] {attachment.filename}",
                f"- content_type: {attachment.content_type}",
                f"- parser: {attachment.parser}",
                f"- size_bytes: {attachment.original_bytes}",
                f"- truncated: {'yes' if attachment.truncated else 'no'}",
            ]
        )
        for warning in attachment.warnings:
            sections.append(f"- warning: {warning}")
        sections.extend(["- content:", "```text", attachment.text, "```"])
    return "\n".join(sections).strip()


def parse_form_json_field(raw: str | None, *, field_name: str, expected_type: type[Any]) -> Any:
    if raw is None or not raw.strip():
        return expected_type()
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise FileIngestError(f"Invalid JSON in form field `{field_name}`.") from exc
    if not isinstance(parsed, expected_type):
        expected_name = getattr(expected_type, "__name__", str(expected_type))
        raise FileIngestError(f"Form field `{field_name}` must be a JSON {expected_name}.")
    return parsed


def parse_form_skill_ids(values: Sequence[str], raw_value: str | None) -> list[str] | None:
    normalized_values = [value.strip() for value in values if value and value.strip()]
    if len(normalized_values) > 1:
        return normalized_values
    if normalized_values:
        raw_value = normalized_values[0]
    if raw_value is None or not raw_value.strip():
        return None
    value = raw_value.strip()
    if value.startswith("["):
        parsed = parse_form_json_field(value, field_name="skill_ids", expected_type=list)
        return [str(item).strip() for item in parsed if str(item).strip()]
    if "," in value:
        return [item.strip() for item in value.split(",") if item.strip()]
    return [value]


def _fit_attachment_to_budget(attachment: ChatAttachment, remaining_chars: int) -> ChatAttachment:
    allowed = min(MAX_PROMPT_CHARS_PER_FILE, max(0, remaining_chars))
    text = attachment.text
    warnings = list(attachment.warnings)
    truncated = attachment.truncated

    if allowed == 0:
        warnings.append("Content omitted because the combined attachment prompt budget was exceeded.")
        return ChatAttachment(
            filename=attachment.filename,
            content_type=attachment.content_type,
            parser=attachment.parser,
            text="[content omitted due to prompt budget]",
            original_bytes=attachment.original_bytes,
            truncated=True,
            warnings=warnings,
        )

    if len(text) > allowed:
        text = text[:allowed].rstrip() + "\n...[truncated]"
        warnings.append("Content truncated to fit the prompt budget.")
        truncated = True

    return ChatAttachment(
        filename=attachment.filename,
        content_type=attachment.content_type,
        parser=attachment.parser,
        text=text,
        original_bytes=attachment.original_bytes,
        truncated=truncated,
        warnings=warnings,
    )


def _looks_like_code_file(suffix: str) -> bool:
    return suffix in TEXT_EXTENSIONS


def _decode_text(raw: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "gbk"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _parse_csv(raw: bytes) -> str:
    decoded = _decode_text(raw)
    rows = list(csv.reader(StringIO(decoded)))
    return _format_table_preview(rows, heading="CSV preview")


def _parse_xlsx(raw: bytes) -> str:
    workbook = load_workbook(filename=BytesIO(raw), read_only=True, data_only=True)
    sections: list[str] = []
    for sheet in workbook.worksheets[:MAX_SHEETS]:
        rows = []
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            rows.append([_stringify_cell(value) for value in row[:MAX_TABLE_COLUMNS]])
            if row_index >= MAX_TABLE_ROWS:
                break
        sections.append(_format_table_preview(rows, heading=f"Sheet: {sheet.title}"))
    return "\n\n".join(section for section in sections if section.strip())


def _parse_xls(raw: bytes) -> str:
    workbook = xlrd.open_workbook(file_contents=raw)
    sections: list[str] = []
    for sheet in workbook.sheets()[:MAX_SHEETS]:
        rows = []
        for row_index in range(min(sheet.nrows, MAX_TABLE_ROWS)):
            row = [_stringify_cell(sheet.cell_value(row_index, col_index)) for col_index in range(min(sheet.ncols, MAX_TABLE_COLUMNS))]
            rows.append(row)
        sections.append(_format_table_preview(rows, heading=f"Sheet: {sheet.name}"))
    return "\n\n".join(section for section in sections if section.strip())


def _parse_docx(raw: bytes) -> str:
    document = Document(BytesIO(raw))
    parts: list[str] = []
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    if paragraphs:
        parts.append("Paragraphs:\n" + "\n".join(paragraphs))
    for table_index, table in enumerate(document.tables[:MAX_DOCX_TABLES], start=1):
        rows = []
        for row in table.rows[:MAX_TABLE_ROWS]:
            rows.append([cell.text.strip() for cell in row.cells[:MAX_TABLE_COLUMNS]])
        parts.append(_format_table_preview(rows, heading=f"Table {table_index}"))
    return "\n\n".join(parts)


def _parse_pdf(raw: bytes) -> str:
    reader = PdfReader(BytesIO(raw))
    parts: list[str] = []
    for page_index, page in enumerate(reader.pages[:MAX_PDF_PAGES], start=1):
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(f"Page {page_index}:\n{text}")
    return "\n\n".join(parts)


def _format_table_preview(rows: Sequence[Sequence[Any]], *, heading: str) -> str:
    if not rows:
        return f"{heading}\n[empty]"
    normalized_rows = [[_stringify_cell(cell) for cell in row] for row in rows]
    width = max(len(row) for row in normalized_rows)
    width = min(width, MAX_TABLE_COLUMNS)
    trimmed = [list(row[:width]) + [""] * max(0, width - len(row[:width])) for row in normalized_rows[:MAX_TABLE_ROWS]]
    if width == 0:
        return f"{heading}\n[empty]"
    column_widths = [
        min(40, max(len(str(row[col_index])) for row in trimmed))
        for col_index in range(width)
    ]
    lines = [heading]
    for row in trimmed:
        formatted = " | ".join(_clip_cell(row[col_index], column_widths[col_index]).ljust(column_widths[col_index]) for col_index in range(width))
        lines.append(formatted.rstrip())
    return "\n".join(lines)


def _clip_cell(value: str, width: int) -> str:
    if len(value) <= width:
        return value
    return value[: max(0, width - 3)].rstrip() + "..."


def _stringify_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        text = f"{value:.6f}".rstrip("0").rstrip(".")
        return text or "0"
    return str(value).strip()
