from __future__ import annotations

import csv
import hashlib
import json
import os
from datetime import date, datetime
from pathlib import Path

definition = {
    "name": "pg-risk-dataset-sync",
    "description": "同步企业风险结构化数据到 PostgreSQL，并保证幂等导入。",
}


DEFAULT_TABLE_NAME = "enterprise_risk_events"
DEFAULT_IMPORT_LOG_TABLE = "enterprise_risk_dataset_import_log"
SUPPORTED_SUFFIXES = {".xlsx", ".xls", ".csv"}

HEADER_MAP = {
    "企业名称": "company_name",
    "企业住所地": "company_address",
    "企业所属区域": "region",
    "企业所在行业": "industry",
    "组织形式": "org_form",
    "所有权性质": "ownership_nature",
    "企业规模": "company_size",
    "案件名称": "case_title",
    "案件大类": "case_category",
    "案件类型（子类1）": "case_type_sub1",
    "案件类型（子类2）": "case_type_sub2",
    "受理日期": "accepted_date",
    "企业诉讼地位（大类）": "litigation_role_major",
    "企业诉讼地位": "litigation_role",
    "部门受案号": "department_case_no",
    "科技企业": "tech_enterprise",
}

BUSINESS_COLUMNS = [
    "company_name",
    "company_address",
    "region",
    "industry",
    "org_form",
    "ownership_nature",
    "tech_enterprise",
    "company_size",
    "case_title",
    "case_category",
    "case_type_sub1",
    "case_type_sub2",
    "department_case_no",
    "litigation_role_major",
    "litigation_role",
    "accepted_date",
]


def execute(params):
    explicit_data_dir = str(params.get("data_dir") or "").strip()
    env_data_dir = str(os.getenv("PG_RISK_DATA_DIR") or "").strip()
    default_data_dir = _default_data_dir()
    data_dir = explicit_data_dir or env_data_dir or default_data_dir
    source_files = _string_list(params.get("source_files"))
    table_name = _safe_identifier(str(params.get("table_name") or DEFAULT_TABLE_NAME).strip() or DEFAULT_TABLE_NAME)
    import_log_table = _safe_identifier(
        str(params.get("import_log_table") or DEFAULT_IMPORT_LOG_TABLE).strip() or DEFAULT_IMPORT_LOG_TABLE
    )
    dry_run = bool(params.get("dry_run", False))
    force_rescan = bool(params.get("force_rescan", False))

    resolved_files = resolve_source_files(data_dir=data_dir, source_files=source_files)
    if (
        not resolved_files
        and not source_files
        and not explicit_data_dir
        and env_data_dir
        and env_data_dir != default_data_dir
    ):
        data_dir = default_data_dir
        resolved_files = resolve_source_files(data_dir=data_dir, source_files=source_files)
    if not resolved_files:
        return {
            "success": False,
            "error_type": "missing_source_files",
            "error": "No source files were found for pg-risk-dataset-sync.",
            "data_dir": data_dir,
            "supported_suffixes": sorted(SUPPORTED_SUFFIXES),
        }

    file_summaries = []
    distinct_fingerprints = set()
    total_rows_read = 0
    scanned_records = []

    for path in resolved_files:
        source_file = path.name
        event_source = detect_event_source(path)
        file_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        records = load_source_records(path, event_source=event_source)
        total_rows_read += len(records)
        for record in records:
            distinct_fingerprints.add(record["row_fingerprint"])
        scanned_records.append(
            {
                "path": path,
                "source_file": source_file,
                "event_source": event_source,
                "file_hash": file_hash,
                "records": records,
            }
        )
        file_summaries.append(
            {
                "source_file": source_file,
                "event_source": event_source,
                "file_hash": file_hash,
                "rows_read": len(records),
                "status": "scanned",
            }
        )

    duplicate_rows_in_input = total_rows_read - len(distinct_fingerprints)

    if dry_run:
        return {
            "success": True,
            "mode": "dry_run",
            "data_dir": data_dir,
            "files_seen": len(scanned_records),
            "rows_read": total_rows_read,
            "distinct_rows": len(distinct_fingerprints),
            "duplicate_rows_in_input": duplicate_rows_in_input,
            "file_summaries": file_summaries,
            "table_name": table_name,
            "import_log_table": import_log_table,
        }

    try:
        conn = _connect_postgres()
    except Exception as exc:
        return {
            "success": False,
            "error_type": "postgres_unavailable",
            "error": f"Failed to connect to PostgreSQL: {exc}",
            "table_name": table_name,
            "import_log_table": import_log_table,
        }

    with conn:
        with conn.cursor() as cur:
            ensure_schema(cur, table_name=table_name, import_log_table=import_log_table)
            existing_rows_before = _fetch_one_int(cur, f"SELECT COUNT(*) FROM {table_name}")
            files_imported = 0
            files_skipped_same_hash = 0
            rows_inserted = 0
            rows_deduped = 0

            for item, summary in zip(scanned_records, file_summaries):
                should_skip_same_hash = existing_rows_before > 0 and not force_rescan and _file_hash_exists(
                    cur,
                    import_log_table=import_log_table,
                    source_file=item["source_file"],
                    file_hash=item["file_hash"],
                )
                if should_skip_same_hash:
                    files_skipped_same_hash += 1
                    summary["status"] = "skipped_same_hash"
                    continue

                inserted_for_file = 0
                for record in item["records"]:
                    inserted = insert_record(cur, table_name=table_name, record=record)
                    inserted_for_file += inserted
                deduped_for_file = len(item["records"]) - inserted_for_file
                rows_inserted += inserted_for_file
                rows_deduped += deduped_for_file
                files_imported += 1
                summary["status"] = "imported"
                summary["rows_inserted"] = inserted_for_file
                summary["rows_deduped"] = deduped_for_file
                upsert_import_log(
                    cur,
                    import_log_table=import_log_table,
                    source_file=item["source_file"],
                    file_hash=item["file_hash"],
                    event_source=item["event_source"],
                    row_count=len(item["records"]),
                    inserted_rows=inserted_for_file,
                    deduped_rows=deduped_for_file,
                )

            existing_rows_after = _fetch_one_int(cur, f"SELECT COUNT(*) FROM {table_name}")
            return {
                "success": True,
                "mode": "sync",
                "data_dir": data_dir,
                "table_name": table_name,
                "import_log_table": import_log_table,
                "files_seen": len(scanned_records),
                "files_imported": files_imported,
                "files_skipped_same_hash": files_skipped_same_hash,
                "rows_read": total_rows_read,
                "rows_inserted": rows_inserted,
                "rows_deduped": rows_deduped,
                "duplicate_rows_in_input": duplicate_rows_in_input,
                "existing_rows_before": existing_rows_before,
                "existing_rows_after": existing_rows_after,
                "dataset_ready": existing_rows_after > 0,
                "file_summaries": file_summaries,
            }


def resolve_source_files(data_dir, source_files):
    if source_files:
        resolved = []
        for raw in source_files:
            candidate = Path(raw)
            if not candidate.is_absolute() and data_dir:
                candidate = Path(data_dir) / candidate
            candidate = candidate.expanduser().resolve()
            if candidate.is_file() and candidate.suffix.lower() in SUPPORTED_SUFFIXES:
                resolved.append(candidate)
        return sorted({path for path in resolved})

    if not data_dir:
        return []
    root = Path(data_dir).expanduser().resolve()
    if not root.is_dir():
        return []
    files = [path for path in root.iterdir() if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES]
    return sorted(files)


def _default_data_dir():
    return str((Path(__file__).resolve().parent / "assets" / "input").resolve())


def detect_event_source(path):
    name = path.name
    if "行政处罚" in name:
        return "administrative_penalty"
    return "judicial_case"


def load_source_records(path, event_source):
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        raw_rows = _read_xlsx_rows(path)
    elif suffix == ".xls":
        raw_rows = _read_xls_rows(path)
    elif suffix == ".csv":
        raw_rows = _read_csv_rows(path)
    else:
        raise ValueError(f"Unsupported file type: {path}")

    records = []
    for row in raw_rows:
        record = normalize_record(row, event_source=event_source, source_file=path.name)
        if record:
            records.append(record)
    return records


def normalize_record(raw_row, event_source, source_file):
    normalized = {column: None for column in BUSINESS_COLUMNS}
    raw_json = {}
    has_any_value = False

    for header, value in raw_row.items():
        key = HEADER_MAP.get(_normalize_text(header))
        raw_json[str(header).strip()] = _json_safe_value(value)
        if key is None:
            continue
        cleaned = normalize_cell(value, field=key)
        if cleaned not in (None, ""):
            has_any_value = True
        normalized[key] = cleaned

    if not has_any_value:
        return None

    normalized["company_name"] = normalized["company_name"] or ""
    fingerprint = build_row_fingerprint(event_source=event_source, normalized=normalized)
    record = {
        "event_source": event_source,
        "source_file": source_file,
        "row_fingerprint": fingerprint,
        "raw_json": raw_json,
    }
    record.update(normalized)
    record["raw_json"] = json.dumps(raw_json, ensure_ascii=False)
    return record


def build_row_fingerprint(event_source, normalized):
    payload = {"event_source": event_source}
    for key in BUSINESS_COLUMNS:
        value = normalized.get(key)
        payload[key] = "" if value is None else str(value)
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def ensure_schema(cur, table_name, import_log_table):
    cur.execute(
        f"""
        CREATE TABLE IF NOT EXISTS {table_name} (
            event_id BIGSERIAL PRIMARY KEY,
            event_source TEXT NOT NULL,
            source_file TEXT NOT NULL,
            company_name TEXT NOT NULL DEFAULT '',
            company_address TEXT,
            region TEXT,
            industry TEXT,
            org_form TEXT,
            ownership_nature TEXT,
            tech_enterprise TEXT,
            company_size TEXT,
            case_title TEXT,
            case_category TEXT,
            case_type_sub1 TEXT,
            case_type_sub2 TEXT,
            department_case_no TEXT,
            litigation_role_major TEXT,
            litigation_role TEXT,
            accepted_date DATE,
            raw_json JSONB NOT NULL DEFAULT '{{}}'::jsonb,
            row_fingerprint TEXT NOT NULL UNIQUE,
            imported_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
        )
        """
    )
    cur.execute(
        f"""
        CREATE TABLE IF NOT EXISTS {import_log_table} (
            import_id BIGSERIAL PRIMARY KEY,
            source_file TEXT NOT NULL,
            file_hash TEXT NOT NULL,
            event_source TEXT NOT NULL,
            row_count INTEGER NOT NULL DEFAULT 0,
            inserted_rows INTEGER NOT NULL DEFAULT 0,
            deduped_rows INTEGER NOT NULL DEFAULT 0,
            imported_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            UNIQUE (source_file, file_hash)
        )
        """
    )
    for ddl in [
        f"CREATE INDEX IF NOT EXISTS idx_{table_name}_event_source ON {table_name} (event_source)",
        f"CREATE INDEX IF NOT EXISTS idx_{table_name}_accepted_date ON {table_name} (accepted_date)",
        f"CREATE INDEX IF NOT EXISTS idx_{table_name}_region ON {table_name} (region)",
        f"CREATE INDEX IF NOT EXISTS idx_{table_name}_industry ON {table_name} (industry)",
        f"CREATE INDEX IF NOT EXISTS idx_{table_name}_company_name ON {table_name} (company_name)",
        f"CREATE INDEX IF NOT EXISTS idx_{table_name}_case_category ON {table_name} (case_category)",
    ]:
        cur.execute(ddl)


def insert_record(cur, table_name, record):
    cur.execute(
        f"""
        INSERT INTO {table_name} (
            event_source,
            source_file,
            company_name,
            company_address,
            region,
            industry,
            org_form,
            ownership_nature,
            tech_enterprise,
            company_size,
            case_title,
            case_category,
            case_type_sub1,
            case_type_sub2,
            department_case_no,
            litigation_role_major,
            litigation_role,
            accepted_date,
            raw_json,
            row_fingerprint
        ) VALUES (
            %(event_source)s,
            %(source_file)s,
            %(company_name)s,
            %(company_address)s,
            %(region)s,
            %(industry)s,
            %(org_form)s,
            %(ownership_nature)s,
            %(tech_enterprise)s,
            %(company_size)s,
            %(case_title)s,
            %(case_category)s,
            %(case_type_sub1)s,
            %(case_type_sub2)s,
            %(department_case_no)s,
            %(litigation_role_major)s,
            %(litigation_role)s,
            %(accepted_date)s,
            %(raw_json)s,
            %(row_fingerprint)s
        )
        ON CONFLICT (row_fingerprint) DO NOTHING
        """,
        record,
    )
    return 1 if cur.rowcount > 0 else 0


def upsert_import_log(cur, import_log_table, source_file, file_hash, event_source, row_count, inserted_rows, deduped_rows):
    cur.execute(
        f"""
        INSERT INTO {import_log_table} (
            source_file,
            file_hash,
            event_source,
            row_count,
            inserted_rows,
            deduped_rows
        ) VALUES (%s, %s, %s, %s, %s, %s)
        ON CONFLICT (source_file, file_hash) DO UPDATE SET
            row_count = EXCLUDED.row_count,
            inserted_rows = EXCLUDED.inserted_rows,
            deduped_rows = EXCLUDED.deduped_rows,
            imported_at = NOW()
        """,
        (source_file, file_hash, event_source, row_count, inserted_rows, deduped_rows),
    )


def _file_hash_exists(cur, import_log_table, source_file, file_hash):
    cur.execute(
        f"SELECT 1 FROM {import_log_table} WHERE source_file = %s AND file_hash = %s LIMIT 1",
        (source_file, file_hash),
    )
    return cur.fetchone() is not None


def _fetch_one_int(cur, sql):
    cur.execute(sql)
    row = cur.fetchone()
    return int(row[0] or 0) if row else 0


def _read_xlsx_rows(path):
    from openpyxl import load_workbook

    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    try:
        rows = []
        for sheet in workbook.worksheets:
            rows.extend(_sheet_to_rows(sheet.iter_rows(values_only=True)))
        return rows
    finally:
        workbook.close()


def _read_xls_rows(path):
    import xlrd

    workbook = xlrd.open_workbook(path)
    rows = []
    for sheet in workbook.sheets():
        if sheet.nrows <= 0:
            continue
        header = None
        for row_index in range(sheet.nrows):
            values = sheet.row_values(row_index)
            if header is None:
                if any(_normalize_text(value) for value in values):
                    header = [str(value).strip() for value in values]
                continue
            if not any(_normalize_text(value) for value in values):
                continue
            rows.append(dict(zip(header, values)))
    return rows


def _read_csv_rows(path):
    rows = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            rows.append(dict(row))
    return rows


def _sheet_to_rows(iterable_rows):
    rows = []
    header = None
    for values in iterable_rows:
        if header is None:
            if any(_normalize_text(value) for value in values):
                header = [str(value).strip() if value is not None else "" for value in values]
            continue
        if not any(_normalize_text(value) for value in values):
            continue
        rows.append(dict(zip(header, values)))
    return rows


def normalize_cell(value, field):
    if value is None:
        return None
    if field == "accepted_date":
        return normalize_date(value)
    if isinstance(value, str):
        cleaned = " ".join(value.strip().split())
        return cleaned or None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None


def normalize_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return None
    text = text.replace("/", "-")
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m", "%Y.%m.%d"):
        try:
            parsed = datetime.strptime(text, fmt)
            if fmt == "%Y-%m":
                return parsed.strftime("%Y-%m-01")
            return parsed.date().isoformat()
        except ValueError:
            continue
    return text


def _json_safe_value(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def _normalize_text(value):
    if value is None:
        return ""
    return str(value).strip()


def _safe_identifier(value):
    text = "".join(char for char in value if char.isalnum() or char == "_")
    if not text:
        raise ValueError("table name is empty after sanitization")
    return text


def _string_list(value):
    if not isinstance(value, list):
        return []
    result = []
    for item in value:
        text = str(item or "").strip()
        if text:
            result.append(text)
    return result


def _connect_postgres():
    import psycopg

    dsn = str(os.getenv("PG_DSN") or "").strip()
    if dsn:
        return psycopg.connect(dsn, connect_timeout=5)
    return psycopg.connect(
        host=os.getenv("PGHOST") or None,
        port=os.getenv("PGPORT") or None,
        dbname=os.getenv("PGDATABASE") or os.getenv("POSTGRES_DB") or None,
        user=os.getenv("PGUSER") or os.getenv("POSTGRES_USER") or None,
        password=os.getenv("PGPASSWORD") or os.getenv("POSTGRES_PASSWORD") or None,
        connect_timeout=5,
    )
