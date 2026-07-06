from __future__ import annotations

import json
from io import BytesIO

import agent_runtime.app as app_module
from agent_runtime.file_ingest import build_attachment_context, parse_file_content
from agent_runtime.models import ChatAttachment, ChatResponse
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook


def test_parse_text_and_csv_files():
    text_attachment = parse_file_content("notes.md", "text/markdown", b"# Title\nhello runtime")
    csv_attachment = parse_file_content("metrics.csv", "text/csv", b"name,value\nspeed,12\n")

    assert text_attachment.parser == "text"
    assert "hello runtime" in text_attachment.text
    assert csv_attachment.parser == "csv"
    assert "CSV preview" in csv_attachment.text
    assert "name" in csv_attachment.text


def test_parse_docx_xlsx_pdf_and_xls(monkeypatch):
    docx_buffer = BytesIO()
    document = Document()
    document.add_heading("Maintenance Note", level=1)
    document.add_paragraph("Machine A temperature exceeded threshold.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "alarm_code"
    table.cell(0, 1).text = "severity"
    table.cell(1, 0).text = "A-100"
    table.cell(1, 1).text = "major"
    document.save(docx_buffer)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MachineStatus"
    sheet.append(["machine", "status", "temperature"])
    sheet.append(["MachineA", "running", 72.5])
    xlsx_buffer = BytesIO()
    workbook.save(xlsx_buffer)

    docx_attachment = parse_file_content(
        "maintenance.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        docx_buffer.getvalue(),
    )
    xlsx_attachment = parse_file_content(
        "machine-status.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        xlsx_buffer.getvalue(),
    )
    pdf_attachment = parse_file_content("report.pdf", "application/pdf", _build_minimal_pdf("hello pdf runtime"))

    class FakeSheet:
        name = "LegacySheet"
        nrows = 2
        ncols = 2

        def cell_value(self, row_index, col_index):
            data = [["machine", "status"], ["MachineB", "idle"]]
            return data[row_index][col_index]

    class FakeWorkbook:
        def sheets(self):
            return [FakeSheet()]

    monkeypatch.setattr("agent_runtime.file_ingest.xlrd.open_workbook", lambda file_contents: FakeWorkbook())
    xls_attachment = parse_file_content("legacy.xls", "application/vnd.ms-excel", b"legacy")

    assert docx_attachment.parser == "docx"
    assert "Maintenance Note" in docx_attachment.text
    assert "alarm_code" in docx_attachment.text
    assert xlsx_attachment.parser == "xlsx"
    assert "Sheet: MachineStatus" in xlsx_attachment.text
    assert pdf_attachment.parser == "pdf"
    assert "hello pdf runtime" in pdf_attachment.text.lower()
    assert xls_attachment.parser == "xls"
    assert "LegacySheet" in xls_attachment.text


def test_build_attachment_context_includes_file_blocks():
    message = build_attachment_context(
        [
            ChatAttachment(
                filename="notes.txt",
                content_type="text/plain",
                parser="text",
                text="line1\nline2",
                original_bytes=11,
            )
        ],
    )

    assert "Parsed attachments for the immediately preceding user request." in message
    assert "[Attachment 1] notes.txt" in message
    assert "```text" in message


def test_chat_endpoint_accepts_multipart_uploads(monkeypatch):
    captured = {}

    async def fake_chat(request):
        captured["request"] = request
        return ChatResponse(
            workspace_id=request.workspace_id or "ws_upload",
            conversation_id="conv_upload",
            agent_id=request.agent_id,
            message="ok",
            status="completed",
            steps=[],
            tool_calls=[],
            model="test-model",
        )

    monkeypatch.setattr(app_module.runtime, "chat", fake_chat)
    client = TestClient(app_module.app)

    response = client.post(
        "/chat",
        data={"message": "请总结附件", "metadata": json.dumps({"source": "ui"})},
        files=[
            ("files", ("notes.md", b"# Notes\nHello upload", "text/markdown")),
            ("files", ("metrics.csv", b"name,value\nspeed,12\n", "text/csv")),
        ],
    )

    assert response.status_code == 200
    request = captured["request"]
    assert request.message == "请总结附件"
    assert request.metadata["source"] == "ui"
    assert [item.filename for item in request.attachments] == ["notes.md", "metrics.csv"]
    assert "Hello upload" in request.attachments[0].text
    assert "CSV preview" in request.attachments[1].text


def _build_minimal_pdf(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 18 Tf 72 720 Td ({escaped}) Tj ET"
    objects = [
        "1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n",
        "2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n",
        (
            "3 0 obj\n"
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            "/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>\n"
            "endobj\n"
        ),
        f"4 0 obj\n<< /Length {len(stream.encode('latin-1'))} >>\nstream\n{stream}\nendstream\nendobj\n",
        "5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n",
    ]

    pdf = b"%PDF-1.4\n"
    offsets = [0]
    for obj in objects:
        offsets.append(len(pdf))
        pdf += obj.encode("latin-1")
    xref_start = len(pdf)
    pdf += f"xref\n0 {len(objects) + 1}\n".encode("latin-1")
    pdf += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        pdf += f"{offset:010d} 00000 n \n".encode("latin-1")
    pdf += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_start}\n%%EOF\n".encode("latin-1")
    )
    return pdf
